import pandas as pd
import re
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

# ── 1. Load ──────────────────────────────────────────────────────────────────
df = pd.read_excel("Customer Call List.xlsx")

# ── 2. Drop useless column ───────────────────────────────────────────────────
df.drop(columns=["Not_Useful_Column"], inplace=True, errors="ignore")

# ── 3. Split Address → Street_Address, State, Zip_code ──────────────────────
def parse_address(addr):
    if pd.isna(addr) or str(addr).strip().lower() in ("n/a", "", "nan"):
        return pd.Series({"Street_Address": None, "State": None, "Zip_code": None})
    addr = str(addr).strip()
    zip_match = re.search(r"\b(\d{5})\b", addr)
    zip_code = zip_match.group(1) if zip_match else None
    if zip_code:
        addr = addr.replace(zip_code, "").strip(", ")
    parts = [p.strip() for p in addr.split(",")]
    street = parts[0] if len(parts) >= 1 else None
    state  = parts[-1] if len(parts) >= 2 else None
    return pd.Series({"Street_Address": street, "State": state, "Zip_code": zip_code})

addr_split = df["Address"].apply(parse_address)
df = pd.concat([df.drop(columns=["Address"]), addr_split], axis=1)

# ── 4. Clean First_Name / Last_Name (strip junk chars) ───────────────────────
def clean_name(val):
    if pd.isna(val): return None
    return re.sub(r"[^a-zA-Z\s\-']", "", str(val)).strip().title()

df["First_Name"] = df["First_Name"].apply(clean_name)
df["Last_Name"]  = df["Last_Name"].apply(clean_name)

# ── 5. Standardise Phone_Number → NNN-NNN-NNNN ──────────────────────────────
def clean_phone(val):
    if pd.isna(val): return None
    digits = re.sub(r"\D", "", str(val))
    if len(digits) == 10:
        return f"{digits[:3]}-{digits[3:6]}-{digits[6:]}"
    return None

df["Phone_Number"] = df["Phone_Number"].apply(clean_phone)

# ── 6. Standardise Yes/No columns ───────────────────────────────────────────
def std_yesno(val):
    if pd.isna(val): return None
    v = str(val).strip().lower()
    if v in ("yes", "y"): return "Yes"
    if v in ("no",  "n"): return "No"
    return None

df["Paying Customer"] = df["Paying Customer"].apply(std_yesno)
df["Do_Not_Contact"]  = df["Do_Not_Contact"].apply(std_yesno)

# ── 7. Drop duplicates ───────────────────────────────────────────────────────
df.drop_duplicates(inplace=True)

# ── 8. Remove Do_Not_Contact = Yes ──────────────────────────────────────────
df = df[df["Do_Not_Contact"] != "Yes"].reset_index(drop=True)

# ── 9. Remove rows with no phone number ──────────────────────────────────────
df = df[df["Phone_Number"].notna()].reset_index(drop=True)

# ── 10. Final column order & rename ─────────────────────────────────────────
df = df[["CustomerID","First_Name","Last_Name","Phone_Number",
         "Paying Customer","Do_Not_Contact","Street_Address","State","Zip_code"]]
df.rename(columns={"Paying Customer": "Paying_Customer"}, inplace=True)

df.to_excel("cloudeAI_code_clean.xlsx", index=False)